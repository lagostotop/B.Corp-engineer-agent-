import csv
import json
import os
from typing import Optional

TEXT_EXTENSIONS={
    ".txt",".md",".py",".js",".ts",".jsx",".tsx",
    ".html",".css",".csv",".json"
}

def _read_text(path:str)->str:
    with open(path,"r",encoding="utf-8",errors="ignore") as f:
        return f.read()

def _read_pdf(path:str)->str:
    from PyPDF2 import PdfReader
    reader=PdfReader(path)
    parts=[]
    for page in reader.pages:
        text=page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)

def _read_docx(path:str)->str:
    import docx
    document=docx.Document(path)
    parts=[p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            values=[cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))

    return "\n\n".join(parts)

def _read_csv(path:str)->str:
    rows=[]
    with open(path,"r",encoding="utf-8",errors="ignore",newline="") as f:
        reader=csv.reader(f)
        for row in reader:
            rows.append(" | ".join(row))
    return "\n".join(rows)

def _read_json(path:str)->str:
    with open(path,"r",encoding="utf-8",errors="ignore") as f:
        data=json.load(f)
    return json.dumps(data,ensure_ascii=False,indent=2)

def extract_text(path:str,filename:Optional[str]=None,max_chars:int=500000)->str:
    if not path or not os.path.exists(path):
        raise FileNotFoundError("File not found.")

    name=filename or os.path.basename(path)
    ext=os.path.splitext(name)[1].lower()

    if ext in TEXT_EXTENSIONS:
        if ext==".csv":
            text=_read_csv(path)
        elif ext==".json":
            text=_read_json(path)
        else:
            text=_read_text(path)

    elif ext==".pdf":
        text=_read_pdf(path)

    elif ext==".docx":
        text=_read_docx(path)

    else:
        raise ValueError(
            f"Text extraction is not supported for {ext or 'unknown'} files."
        )

    return str(text or "")[:max_chars].strip()